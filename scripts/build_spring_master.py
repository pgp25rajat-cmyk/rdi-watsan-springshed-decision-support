"""Build the final comprehensive spring master table (one row per common-file record).

Inputs:  data/01_staged/common/csv/spring_records_common_v1.csv
         data/02_processed/spring_features/spring_features_v1.csv
         data/02_processed/spring_features/spring_secondary_features_v1.csv (optional; scripts/build_spring_secondary_features.py)
Output:  data/02_processed/spring_master/spring_master_v1.csv (+ data dictionary,
         fill summary, duplicate pairs, and one-row-per-spring files spring_master_v1_clean.csv,
         _input.csv and _test.csv) and spring_master_v1.xlsx

How gaps are filled (nothing is invented or interpolated; every filled value says
where it came from in a *_from column):
- Terrain, land cover, rainfall, soil and lineaments are mapped layers, so they are
  given for every spring with a usable coordinate, from the same layer. The source's
  own value is kept beside it for comparison.
- Field-only values (discharge, pH, TDS, spring type, lithology, recharge area) come
  from the record itself. Where the record lacks one, a record from another source
  confirmed as the same spring (scripts/spring_duplicates.py) fills it, named in the *_from
  column. Otherwise the value stays blank and gaps_remaining lists it.
- Duplicates: records within 30 m are sorted into same spring / different springs / needs
  review. Team decisions in data/01_staged/common/duplicate_review_decisions_v1.csv override
  the rules on the next build.
Usage: python scripts/build_spring_master.py
"""

import csv
import math
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from build_spring_secondary_features import COLUMNS as SECONDARY_COLUMNS, OUT_CSV as SECONDARY_CSV
from field_visit_spring_records import field_flows
from spring_duplicates import load_decisions, resolve_group


ROOT = Path(__file__).resolve().parents[1]
COMMON_CSV = ROOT / "data" / "01_staged" / "common" / "csv" / "spring_records_common_v1.csv"
FEATURES_CSV = ROOT / "data" / "02_processed" / "spring_features" / "spring_features_v1.csv"
WATSAN_DISCHARGE_CSV = ROOT / "data" / "01_staged" / "watsan" / "csv" / "source_sheets" / "watsan_discharge_lpm_source_v1.csv"
WORLDCOVER = {"10": "Tree cover", "20": "Shrubland", "30": "Grassland", "40": "Cropland", "50": "Built-up",
              "60": "Bare / sparse vegetation", "80": "Permanent water bodies"}
DISTRICTS = ["Almora", "Bageshwar", "Chamoli", "Champawat", "Dehradun", "Haridwar", "Nainital", "Pauri Garhwal", "Pithoragarh",
             "Rudraprayag", "Tehri Garhwal", "Udham Singh Nagar", "Uttarkashi"]
OUT_DIR = ROOT / "data" / "02_processed" / "spring_master"
DECISIONS_CSV = ROOT / "data" / "01_staged" / "common" / "duplicate_review_decisions_v1.csv"
# Test set, kept out of model building and tuning (agreed with Rajat 2026-10-14): one whole district, the
# Spring Atlas profile springs (the only real recharge-area delineations with coordinates), every record in
# the same 30 m group as a test record, and every spring within TEST_BUFFER_M of one, so no test spring
# has a near-identical neighbour in the input data.
TEST_DISTRICTS = {"Tehri Garhwal"}
TEST_SOURCE_PREFIXES = ("ATP",)
TEST_BUFFER_M = 2000
DECISION_COLUMNS = ["record_id_a", "record_id_b", "decision", "decided_by", "decided_on", "note"]

# Preferred record when several sources list the same spring: richest field data first.
SOURCE_PRIORITY = {"ATP": 0, "FLD": 1, "RDI": 1, "HIM": 2, "WAT": 3, "DAN": 4, "RAT": 5, "ATL": 6}

SPRING_TYPE_CLASSES = [
    ("karst", "Karst"), ("fractur", "Fracture"), ("fault", "Fracture"), ("joint", "Fracture"),
    ("contact", "Contact"), ("contect", "Contact"), ("depression", "Depression"),
    ("deperession", "Depression"), ("deppression", "Depression"), ("dipression", "Depression"), ("seepage", "Seepage"),
    ("gravity", "Gravity / seepage"), ("weathered", "Fracture"), ("stream", "Stream (not a spring)"),
]
LITHOLOGY_CLASSES = [
    ("dolomit", "Carbonate"), ("limestone", "Carbonate"), ("quartzit", "Quartzite"), ("quarzit", "Quartzite"), ("quatzit", "Quartzite"), ("qaurtzit", "Quartzite"),
    ("phyllit", "Phyllite"), ("schist", "Schist"), ("slate", "Slate"), ("shale", "Shale"),
    ("sandstone", "Sandstone"), ("gneiss", "Gneiss / granite"), ("granite", "Gneiss / granite"),
    ("conglomerate", "Conglomerate"), ("alluvi", "Alluvium / debris"), ("unconsolidated", "Alluvium / debris"), ("metabasic", "Metabasic"), ("debris", "Alluvium / debris"),
    ("meta-volcanic", "Metabasic"), ("metavolcanic", "Metabasic"), ("basic", "Metabasic"), ("gravel", "Alluvium / debris"),
    ("boulder", "Alluvium / debris"), ("sand, silt", "Alluvium / debris"), ("silt, clay", "Alluvium / debris"), ("silt,clay", "Alluvium / debris"),
    ("mudstone", "Shale"), ("claystone", "Shale"), ("siltstone", "Sandstone"), ("marble", "Carbonate"), ("carbonate", "Carbonate"), ("amphibolite", "Metabasic"),
]

FIELD_ONLY = ["discharge_lpm_low", "discharge_lpm_high", "ph", "tds_mg_l", "spring_type_source_text",
              "lithology_source_text", "recharge_area_ha_low", "recharge_area_ha_high"]

# (column, unit, from, meaning)
COLUMNS = [
    ("record_id", "", "Common file", "Record id; prefix shows the source (WAT, HIM, ATL, ATP, RDI, RAT, DAN, FLD = field visit)."),
    ("source_dataset", "", "Common file", "Dataset the record came from."),
    ("spring_name", "", "Common file", "Spring name as given by the source."),
    ("village", "", "Common file", ""),
    ("block", "", "Common file", ""),
    ("district", "", "Processed features", "District with spellings unified."),
    ("area_name", "", "Common file", "Area / watershed label (WATSAN, literature)."),
    ("latitude", "decimal degrees", "Processed features", ""),
    ("longitude", "decimal degrees", "Processed features", ""),
    ("coordinate_status", "", "Processed features", "ok, missing or outside_study_region. Mapped values exist only for ok."),
    ("coordinate_precision_decimals", "", "Processed features", "2 or fewer decimals means about 1 km uncertainty."),
    ("coordinate_correction", "", "Processed features", "What was corrected in the source coordinate and why. Only 2 records, where the fix is certain."),
    ("colocated_group_id", "", "Processed features", "Records within 30 m of each other; probably the same spring listed by two sources."),
    ("colocated_group_size", "", "Processed features", ""),
    ("is_primary_record", "", "Derived", "Y for one record per spring: the only record, or the richest record among those confirmed as the same spring. Use Y rows to count distinct springs."),
    ("spring_id", "", "Derived", "Record id of the primary record of this spring. Records with the same spring_id are one spring listed more than once."),
    ("dataset_split", "", "Derived", "input = use to build and tune the tool; test = keep aside until the tool is finished, then score once. Test: Tehri Garhwal, the 10 Spring Atlas profile springs, their 30 m groups, and springs within 2 km of these."),
    ("split_reason", "", "Derived", "Why a record is in the test set."),
    ("duplicate_status", "", "Derived", "Blank for records with no other record within 30 m. Otherwise: same spring, different springs, partly the same spring, or needs review (for the whole 30 m group)."),
    ("duplicate_group", "", "Derived", "For records within 30 m of another: the group, which records are the same spring, which are separate springs and why, and which still need a person's decision. Coloured in the workbook; every pair is on the Duplicate pairs sheet."),
    # terrain
    ("elevation_m", "m", "Copernicus DEM GLO-30", "DEM elevation at the spring (same layer for all springs)."),
    ("elevation_source_m", "m", "Common file", "Elevation the source gave, for comparison."),
    ("elevation_check", "", "Derived", "Only where DEM and source elevation differ by more than 150 m: which is probably wrong. Copied coordinate (another record at the same point fits the DEM), coordinate a little off (the source elevation is found within 300 m), coordinate off by the stated distance, or no point within 2 km has the source elevation."),
    ("slope_deg", "degrees", "Copernicus DEM GLO-30", "Slope at the spring."),
    ("slope_mean_500m_deg", "degrees", "Copernicus DEM GLO-30", "Mean slope within 500 m."),
    ("aspect_class", "", "Copernicus DEM GLO-30", "8-point aspect or Flat."),
    ("relief_1km_m", "m", "Copernicus DEM GLO-30", "Highest minus lowest elevation within 1 km."),
    ("upslope_relief_1km_m", "m", "Copernicus DEM GLO-30", "Highest elevation within 1 km minus the spring elevation."),
    ("tpi_300m_m", "m", "Copernicus DEM GLO-30", "Negative = hollow / valley position."),
    ("drainage_area_ha", "ha", "Copernicus DEM GLO-30 (derived)", "Surface area draining to the spring cell (multiple-flow-direction routing on the filled DEM). The topographic catchment above the spring; not the recharge area, which follows the rocks."),
    ("drainage_area_max30m_ha", "ha", "Copernicus DEM GLO-30 (derived)", "Largest drainage area within about 30 m (allows for GPS error). Much larger than drainage_area_ha means a hollow or stream line is right next to the point."),
    ("wetness_index_twi", "", "Copernicus DEM GLO-30 (derived)", "Topographic wetness index ln(a / tan slope). Higher = wetter, more convergent spot. Hillslopes ~4-8, hollows and valley floors ~8-14."),
    ("distance_to_stream_m", "m", "Copernicus DEM GLO-30 (derived)", "Straight-line distance to the nearest DEM stream (cells draining at least 0.5 km2)."),
    ("nearest_stream_order", "", "Copernicus DEM GLO-30 (derived)", "Strahler order of that stream (1 = headwater)."),
    ("height_above_stream_m", "m", "Copernicus DEM GLO-30 (derived)", "Height above the stream the spring drains to (HAND)."),
    ("slope_source_deg", "degrees", "Common file", "Slope the source gave, for comparison."),
    # land cover
    ("land_cover_point", "", "ESA WorldCover 2021", "Class at the spring (10 m)."),
    ("land_cover_majority_500m", "", "ESA WorldCover 2021", "Most common class within 500 m."),
    ("tree_cover_500m_pct", "%", "ESA WorldCover 2021", ""),
    ("cropland_500m_pct", "%", "ESA WorldCover 2021", ""),
    ("built_up_500m_pct", "%", "ESA WorldCover 2021", ""),
    ("buildings_1km_2016", "count", "Google Open Buildings 2.5D Temporal", "Estimated buildings within 1 km in 2016."),
    ("buildings_1km_2023", "count", "Google Open Buildings 2.5D Temporal", "Estimated buildings within 1 km in 2023."),
    ("building_growth_1km_pct", "%", "Derived", "Buildings added 2016-2023 as a share of the 2016 count (base at least 25). A stand-in for population growth near the spring."),
    ("land_use_source_text", "", "Common file", "Land use as the source described it. WATSAN gives ESA WorldCover codes, decoded here; the class names in WATSAN's own sheet are wrong (e.g. 10 labelled Built-up)."),
    # climate
    ("rainfall_mean_annual_mm", "mm/year", "CHIRPS v2 2018-2025 (FAO WaPOR)", "Mean annual rainfall, ~5 km."),
    ("rainfall_cv_pct", "%", "CHIRPS v2 2018-2025", "Year-to-year variation."),
    ("rainfall_monsoon_share_pct", "%", "CHIRPS v2 2018-2025", "Share falling June-September."),
    ("rainfall_dry_season_mm", "mm", "CHIRPS v2 2018-2025", "November-April rain."),
    ("et0_mean_annual_mm", "mm/year", "FAO WaPOR v3", "Reference evapotranspiration."),
    ("aridity_index", "", "Derived", "Rainfall / ET0."),
    ("rainfall_imd_mean_annual_mm", "mm/year", "IMD 0.25 degree grid, 2001-2025 (2012 missing)", "Mean annual rainfall, ~25 km. Cross-check on CHIRPS; IMD reads lower in the hills."),
    ("rainfall_imd_min_annual_mm", "mm/year", "IMD 0.25 degree grid, 2001-2025", "Driest year."),
    ("rainfall_imd_cv_pct", "%", "IMD 0.25 degree grid, 2001-2025", "Year-to-year variation."),
    ("rainfall_imd_max_1day_mm", "mm", "IMD 0.25 degree grid, 2001-2025", "Wettest single day."),
    ("rainfall_imd_trend_pct_per_decade", "% per decade", "IMD 0.25 degree grid, 2001-2025", "Theil-Sen trend of yearly rain as a share of the mean, per 10 years; negative = getting drier."),
    ("rainfall_imd_trend_p", "", "IMD 0.25 degree grid, 2001-2025", "Mann-Kendall p-value; below 0.05 = a trend unlikely to be chance."),
    ("rainfall_imd_dry_season_trend_pct_per_decade", "% per decade", "IMD 0.25 degree grid, 2001-2025", "Same for January-April plus November-December rain."),
    ("rainfall_source_mm", "mm/year", "Common file", "Rainfall the source gave, for comparison."),
    # soil
    ("soil_clay_pct", "%", "SoilGrids 2.0, 0-30 cm (HWSD v2 0-20 cm where missing)", "Depth-weighted mean of the 0-5, 5-15 and 15-30 cm layers."),
    ("soil_silt_pct", "%", "SoilGrids 2.0, 0-30 cm (HWSD v2 where missing)", ""),
    ("soil_sand_pct", "%", "SoilGrids 2.0, 0-30 cm (HWSD v2 where missing)", ""),
    ("soil_texture_from", "", "Derived", "Which layer gave clay/silt/sand."),
    ("soil_coarse_fragments_pct", "%", "SoilGrids 2.0, 0-30 cm", "Stones and gravel by volume. More coarse fragments = faster infiltration, less storage."),
    ("soil_bulk_density_g_cm3", "g/cm3", "SoilGrids 2.0, 0-30 cm", ""),
    ("soil_org_carbon_density_kg_m3", "kg/m3", "SoilGrids 2.0, 0-30 cm", ""),
    ("soil_type_wrb", "", "HWSD v2.0", "Dominant soil type (~1 km)."),
    ("soil_drainage", "", "HWSD v2.0", ""),
    ("soil_awc_mm_per_m", "mm/m", "HWSD v2.0", "Available water capacity."),
    # geology
    ("lithology_source_text", "", "Common file / colocated record", "Rock type as described in the field (Himmotthan, Atlas profiles, Chamba report, WATSAN rock type)."),
    ("lithology_class", "", "Derived", "Main rock class from the text: Carbonate, Quartzite, Phyllite, Schist, Slate, ... or Mixed."),
    ("lithology_from", "", "Derived", "Record the lithology text came from."),
    ("geological_structure_text", "", "Common file", "Structure notes (WATSAN, Atlas profiles, Chamba report)."),
    ("lineament_distance_m", "m", "DEM lineament candidates", "Distance to nearest candidate lineament. Unvalidated."),
    ("lineament_density_1km", "km/km2", "DEM lineament candidates", ""),
    ("gsi_lithology", "", "GSI NGDR lithology (Bhukosh)", "Rock description of the GSI map polygon at the spring. Uttarakhand only."),
    ("gsi_lithology_class", "", "Derived", "Main rock class from the GSI description, same classes as lithology_class."),
    ("gsi_group", "", "GSI NGDR lithology", "Stratigraphic group (Garhwal, Jaunsar, Central Crystalline, ...)."),
    ("gsi_formation", "", "GSI NGDR lithology", ""),
    ("gsi_age", "", "GSI NGDR lithology", ""),
    ("gsi_distance_to_contact_m", "m", "GSI NGDR lithology", "Distance to the edge of the spring's map polygon (rock contact)."),
    ("rock_class", "", "Derived", "Best rock class: field description (lithology_class) where there is one, else the GSI map (gsi_lithology_class)."),
    ("rock_class_from", "", "Derived", "field record, or GSI map."),
    ("gsi_fault_distance_m", "m", "GSI NGDR faults (Bhukosh)", "Distance to the nearest mapped fault or thrust."),
    ("gsi_nearest_fault_type", "", "GSI NGDR faults", ""),
    ("gsi_nearest_fault_name", "", "GSI NGDR faults", "Named where GSI names it (MCT, MBT, North Almora Thrust, ...)."),
    ("gsi_thrust_distance_m", "m", "GSI NGDR faults", "Distance to the nearest mapped thrust."),
    ("gsi_fault_density_1km", "km/km2", "GSI NGDR faults", "Mapped fault length within 1 km, per km2."),
    ("gsi_fold_distance_m", "m", "GSI NGDR folds (Bhukosh)", "Distance to the nearest mapped fold axis."),
    ("gsi_nearest_fold_type", "", "GSI NGDR folds", "Anticline, syncline, antiform, synform or overturned."),
    ("gsi_anticline_distance_m", "m", "GSI NGDR folds", "Nearest anticline or antiform axis (fractured crest, rocks dip away)."),
    ("gsi_syncline_distance_m", "m", "GSI NGDR folds", "Nearest syncline or synform axis (rocks dip inward and can hold water)."),
    # spring and water
    ("spring_type_source_text", "", "Common file / colocated record", ""),
    ("spring_type_class", "", "Derived", "Fracture, Contact, Depression, Karst, Seepage, Gravity / seepage, or Stream (not a spring)."),
    ("spring_type_from", "", "Derived", ""),
    ("discharge_lpm_low", "L/min", "Common file / colocated record", "Lowest number in the source discharge (Rathi: max flow converted from m3/day)."),
    ("discharge_lpm_high", "L/min", "Common file / colocated record", "Highest number in the source discharge."),
    ("discharge_text", "", "Common file / colocated record", "Discharge as the source wrote it (season, dates)."),
    ("discharge_dry_lpm", "L/min", "WATSAN discharge sheet", "Dry-season figure WATSAN derived from its text estimate (WATSAN springs only)."),
    ("discharge_monsoon_lpm", "L/min", "WATSAN discharge sheet", "Monsoon figure WATSAN derived from its text estimate (WATSAN springs only)."),
    ("discharge_from", "", "Derived", ""),
    ("ph", "", "Common file / colocated record", ""),
    ("tds_mg_l", "mg/l", "Common file / colocated record", ""),
    ("water_quality_from", "", "Derived", ""),
    ("recharge_area_ha_low", "ha", "Common file / colocated record", "WATSAN estimate, Atlas field delineation, Rathi springshed (km2 x 100)."),
    ("recharge_area_ha_high", "ha", "Common file / colocated record", ""),
    ("recharge_area_from", "", "Derived", ""),
    ("seasonality_text", "", "Common file", "WATSAN seasonality label."),
    ("intervention_text", "", "Common file", "Recharge or DI works reported (Atlas profiles, Chamba, Daniel)."),
    ("vulnerability_index_source", "", "Common file", "Daniel 2021 index (one spring)."),
    # bookkeeping
    ("key_fields_filled", "of 13", "Derived", "How many of 13 key fields have a value: elevation, slope, land cover, rainfall, soil texture, rock type, fault distance, lineament distance, spring type, discharge, pH, TDS, recharge area."),
    ("gaps_remaining", "", "Derived", "Key fields still blank and why."),
    ("notes", "", "Processed features / Derived", "Coordinate and elevation notes, values removed as invalid, and values to check."),
]
# Secondary layers collected in October 2026 (data columns only; no module reads them yet) go after the GSI structure columns
_at = [c[0] for c in COLUMNS].index("spring_type_source_text")
COLUMNS[_at:_at] = SECONDARY_COLUMNS
NAMES = [c[0] for c in COLUMNS]


def read(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def classify(text, table, mixed_label=None):
    lowered = (text or "").lower()
    if not lowered.strip():
        return ""
    hits = []
    for key, label in table:
        if key in lowered and label not in hits:
            hits.append(label)
    if not hits:
        return "Other"
    if mixed_label and len(hits) > 1:
        return f"{mixed_label} ({', '.join(hits)})"
    return hits[0]


def field_values(common, feature):
    """Field-only values carried by this record itself."""
    lith = common["lithology_source_value"].strip() or common["rock_type_source_value"].strip()
    max_flow = common["discharge_max_m3_per_day_source_value"]
    max_flow = f"{max_flow} m3/day max flow" if max_flow else ""
    values = {
        "discharge_lpm_low": feature["source_discharge_lpm_low"] or feature["source_max_flow_lpm_converted"],
        "discharge_lpm_high": feature["source_discharge_lpm_high"] or feature["source_max_flow_lpm_converted"],
        "discharge_text": common["discharge_lpm_source_value"] or max_flow,
        # pH 0 and TDS 0 are blanks coded as zero in the Himmotthan layer, not measurements
        "ph": feature["source_ph"] if feature["source_ph"] and 3 <= float(feature["source_ph"]) <= 11 else "",
        "tds_mg_l": feature["source_tds_mg_l"] if feature["source_tds_mg_l"] and float(feature["source_tds_mg_l"]) > 0 else "",
        "spring_type_source_text": common["spring_type_source_value"].strip(),
        "lithology_source_text": lith,
        "recharge_area_ha_low": feature["source_recharge_area_ha_low"] or feature["source_springshed_area_ha_converted"],
        "recharge_area_ha_high": feature["source_recharge_area_ha_high"] or feature["source_springshed_area_ha_converted"],
    }
    return {k: (v or "") for k, v in values.items()}


def land_use_text(record_id, text):
    """WATSAN lulc values are ESA WorldCover class codes (they match the WorldCover pixel for most springs);
    the class names in WATSAN's own lulc sheet do not fit the codes, so they are not used."""
    if record_id.startswith("WAT") and text in WORLDCOVER:
        return f"{text} = {WORLDCOVER[text]} (ESA WorldCover code)"
    return text


def district_from_area(text):
    """District named in a WATSAN area label such as 'Palyalgaon, Yamkeshwar, Pauri Garhwal, Uttarakhand, 249303, India'."""
    parts = [p.strip().lower() for p in (text or "").split(",")]
    for name in DISTRICTS:
        if name.lower() in parts or name.split()[0].lower() in parts:
            return name
    return ""


def richness(record_id, values, field_verified=frozenset()):
    """Preferred record of a spring: one checked on a field visit, then the richest field data, then the source."""
    return (record_id in field_verified, sum(1 for k in FIELD_ONLY if values[k]), -SOURCE_PRIORITY.get(record_id[:3], 9))


def assign_split(rows, groups):
    by_id = {r["record_id"]: r for r in rows}
    reason = {}
    for r in rows:
        if r["district"] in TEST_DISTRICTS:
            reason[r["record_id"]] = f"{r['district']} (held-out district)"
        elif r["record_id"].startswith(TEST_SOURCE_PREFIXES):
            reason[r["record_id"]] = "Spring Atlas profile spring (delineated recharge area)"
    for members in groups.values():
        held = [m for m in members if m in reason]
        for m in members:
            if held and m not in reason:
                reason[m] = f"same 30 m group as {held[0]}"
    anchors = [(by_id[rid]["latitude"], by_id[rid]["longitude"], rid) for rid in reason if by_id[rid]["latitude"]]
    for r in rows:
        if r["record_id"] in reason or not r["latitude"]:
            continue
        lat, lon = float(r["latitude"]), float(r["longitude"])
        for a_lat, a_lon, rid in anchors:
            d = math.hypot((lat - float(a_lat)) * 111_000, (lon - float(a_lon)) * 111_000 * math.cos(math.radians(lat)))
            if d <= TEST_BUFFER_M:
                reason[r["record_id"]] = f"within {d / 1000:.1f} km of test spring {rid}"
                break
    for r in rows:
        r["dataset_split"] = "test" if r["record_id"] in reason else "input"
        r["split_reason"] = reason.get(r["record_id"], "")


def duplicate_label(rid, group, size, spring_of, pairs, team_linked=False):
    if not group:
        return ""
    parts = [f"{group} ({size} records " + ("within 30 m or joined by a team decision)" if team_linked else "within 30 m)")]
    same = [b if a == rid else a for a, b, d, _ in pairs if rid in (a, b) and d == "same"]
    if same:
        parts.append(f"same spring as {', '.join(same)}; " + ("this is the primary record" if spring_of[rid] == rid else f"primary is {spring_of[rid]}"))
    for a, b, d, reason in pairs:
        if rid in (a, b) and d != "same":
            other = b if a == rid else a
            parts.append(("separate spring from " if d == "different" else "needs review against ") + f"{other} ({reason})")
    return "; ".join(parts)


def elevation_check(f, shared_with, features_by_id):
    gap = f["dem_minus_source_elevation_m"]
    if not gap or abs(float(gap)) <= 150:
        return ""
    for other in shared_with:
        o = features_by_id[other]
        if o["dem_minus_source_elevation_m"] and abs(float(o["dem_minus_source_elevation_m"])) <= 50:
            return f"coordinate probably copied: {other} has the same point and its elevation fits the DEM"
    rounded = " (coordinate rounded to about 1 km)" if f["coordinate_precision_decimals"] and int(f["coordinate_precision_decimals"]) <= 2 else ""
    match = f["elevation_match_distance_m"]
    if not match:
        return "no point within 2 km has the source elevation: coordinate or source elevation badly wrong" + rounded
    match = float(match)
    if match <= 300:
        return f"coordinate a little off: the source elevation is {match:.0f} m away on the slope" + rounded
    return f"coordinate off by about {match / 1000:.1f} km (nearest point at the source elevation)" + rounded


def main():
    common = {r["record_id"]: r for r in read(COMMON_CSV)}
    watsan_seasonal = {(r["Latitude"], r["Longitude"]): r for r in read(WATSAN_DISCHARGE_CSV)}
    field_flow = field_flows()
    features = read(FEATURES_CSV)
    secondary = {r["record_id"]: r for r in read(SECONDARY_CSV)} if SECONDARY_CSV.exists() else {}
    own = {f["record_id"]: field_values(common[f["record_id"]], f) for f in features}
    features_by_id = {f["record_id"]: f for f in features}
    report_flow = {rid: own[rid]["discharge_text"] for rid in field_flow if rid in own}
    for rid in report_flow:  # field form flow first, so records of the same spring are filled with it too
        low, high, text = field_flow[rid]
        own[rid].update(discharge_lpm_low=low, discharge_lpm_high=high, discharge_text=text)
    field_verified = frozenset(f["record_id"] for f in features
                               if f["record_id"].startswith("FLD-") or f["coordinate_correction"].startswith("field GPS"))
    groups = {}
    for f in features:
        if f["colocated_group_id"]:
            groups.setdefault(f["colocated_group_id"], []).append(f["record_id"])
    links = {frozenset((rid, c["candidate_match_record_ids"].split()[0])) for rid, c in common.items()
             if "matching name/coordinates" in c["candidate_match_record_ids"]}
    decisions = load_decisions(DECISIONS_CSV)
    # A team "same" decision also joins records further apart than 30 m (e.g. a field GPS against an older survey point).
    group_of = {rid: g for g, members in groups.items() for rid in members}
    team_linked = set()
    for key, (decision, _) in sorted(decisions.items(), key=lambda item: sorted(item[0])):
        a, b = sorted(key)
        if decision != "same" or a not in features_by_id or b not in features_by_id:
            continue
        ga, gb = group_of.get(a), group_of.get(b)
        if ga and ga == gb:
            continue
        target = ga or gb or f"LINK-{sum(1 for g in groups if g.startswith('LINK-')) + 1:04d}"
        groups.setdefault(target, [])
        for other in {ga, gb} - {target, None}:
            groups[target].extend(groups.pop(other))
        for rid in (a, b):
            if rid not in groups[target]:
                groups[target].append(rid)
        for rid in groups[target]:
            group_of[rid] = target
        team_linked.add(target)
    spring_of, cluster_of, status_of, group_pairs = {}, {}, {}, {}
    for group, members in groups.items():
        records = [{"record_id": m, "spring_name": common[m]["spring_name"], "village": common[m]["village"],
                    "elevation_source_m": features_by_id[m]["source_elevation_m"]} for m in members]
        pairs, clusters, status = resolve_group(records, links, decisions)
        group_pairs[group] = pairs
        for cluster in clusters:
            main_record = max(cluster, key=lambda rid: richness(rid, own[rid], field_verified))
            for rid in cluster:
                spring_of[rid], cluster_of[rid], status_of[rid] = main_record, cluster, status
    at_point = {}
    for f in features:
        if f["coordinate_status"] == "ok":
            at_point.setdefault((f["latitude_dd"], f["longitude_dd"]), []).append(f["record_id"])

    rows = []
    for f in features:
        rid = f["record_id"]
        c = common[rid]
        partners = [p for p in cluster_of.get(rid, []) if p != rid]
        partners.sort(key=lambda p: richness(p, own[p], field_verified), reverse=True)

        def take(keys):
            """Fill a set of keys together from this record, else the richest partner that has them."""
            if any(own[rid][k] for k in keys):
                return {k: own[rid][k] for k in keys}, "this record"
            for p in partners:
                if any(own[p][k] for k in keys):
                    near = f["colocated_group_id"] and features_by_id[p]["colocated_group_id"] == f["colocated_group_id"]
                    return {k: own[p][k] for k in keys}, f"colocated record {p} (same spring, {'within 30 m' if near else 'team decision'})"
            return {k: "" for k in keys}, ""

        row = {n: "" for n in NAMES}
        row.update({
            "record_id": rid, "source_dataset": f["source_dataset"], "spring_name": f["spring_name"], "village": f["village"],
            "block": c["block"], "district": f["district_clean"] or district_from_area(c["source_area_name"]),
            "area_name": c["source_area_name"],
            "latitude": f["latitude_dd"], "longitude": f["longitude_dd"], "coordinate_status": f["coordinate_status"],
            "coordinate_precision_decimals": f["coordinate_precision_decimals"],
            "colocated_group_id": f["colocated_group_id"], "colocated_group_size": f["colocated_group_size"],
            "coordinate_correction": f["coordinate_correction"],
            "is_primary_record": "Y" if spring_of.get(rid, rid) == rid else "N",
            "spring_id": spring_of.get(rid, rid),
            "duplicate_status": status_of.get(rid, ""),
            "duplicate_group": duplicate_label(rid, group_of.get(rid, ""), len(groups.get(group_of.get(rid), [])), spring_of,
                                               group_pairs.get(group_of.get(rid), []), group_of.get(rid) in team_linked),
            "drainage_area_ha": f["dem_drainage_area_ha"], "drainage_area_max30m_ha": f["dem_drainage_area_max30m_ha"],
            "wetness_index_twi": f["dem_twi"], "distance_to_stream_m": f["dem_distance_to_stream_m"],
            "nearest_stream_order": f["dem_nearest_stream_order"], "height_above_stream_m": f["dem_height_above_stream_m"],
            "gsi_fold_distance_m": f["distance_to_gsi_fold_axis_m"], "gsi_nearest_fold_type": f["gsi_nearest_fold_type"],
            "gsi_anticline_distance_m": f["distance_to_gsi_anticline_m"], "gsi_syncline_distance_m": f["distance_to_gsi_syncline_m"],
            "elevation_m": f["dem_elevation_m"], "elevation_source_m": f["source_elevation_m"],
            "elevation_check": elevation_check(f, [o for o in at_point.get((f["latitude_dd"], f["longitude_dd"]), []) if o != rid], features_by_id),
            "slope_deg": f["dem_slope_deg"], "slope_mean_500m_deg": f["dem_mean_slope_500m_deg"], "aspect_class": f["dem_aspect_class"],
            "relief_1km_m": f["dem_relief_1km_m"], "upslope_relief_1km_m": f["dem_upslope_relief_1km_m"], "tpi_300m_m": f["dem_tpi_300m_m"],
            "slope_source_deg": f["source_slope_deg"],
            "land_cover_point": f["wc_point_class"], "land_cover_majority_500m": f["wc_majority_500m_class"],
            "tree_cover_500m_pct": f["wc_share_500m_tree_pct"], "cropland_500m_pct": f["wc_share_500m_cropland_pct"],
            "built_up_500m_pct": f["wc_share_500m_built_pct"],
            "buildings_1km_2016": f["obt_buildings_1km_2016"], "buildings_1km_2023": f["obt_buildings_1km_2023"],
            "building_growth_1km_pct": f["obt_building_growth_pct"], "land_use_source_text": land_use_text(rid, c["land_use_source_value"]),
            "rainfall_mean_annual_mm": f["rain_chirps_mean_annual_mm_2018_2025"], "rainfall_cv_pct": f["rain_chirps_cv_annual_pct"],
            "rainfall_monsoon_share_pct": f["rain_chirps_monsoon_share_pct"], "rainfall_dry_season_mm": f["rain_chirps_dry_season_mm"],
            "et0_mean_annual_mm": f["et0_wapor_mean_annual_mm_2018_2025"], "aridity_index": f["aridity_index_p_over_et0"],
            "rainfall_imd_mean_annual_mm": f["rain_imd_mean_annual_mm"], "rainfall_imd_min_annual_mm": f["rain_imd_min_annual_mm"],
            "rainfall_imd_cv_pct": f["rain_imd_cv_annual_pct"], "rainfall_imd_max_1day_mm": f["rain_imd_max_1day_mm"],
            "rainfall_imd_trend_pct_per_decade": f["rain_imd_trend_pct_per_decade"], "rainfall_imd_trend_p": f["rain_imd_trend_kendall_p"],
            "rainfall_imd_dry_season_trend_pct_per_decade": f["rain_imd_dry_season_trend_pct_per_decade"],
            "rainfall_source_mm": f["source_rainfall_mm"],
            "soil_coarse_fragments_pct": f["sg_coarse_fragments_0_30cm_pct"],
            "soil_bulk_density_g_cm3": f["sg_bulk_density_0_30cm_g_cm3"], "soil_org_carbon_density_kg_m3": f["sg_org_carbon_density_0_30cm_kg_m3"],
            "soil_type_wrb": f["soil_dominant_wrb"], "soil_drainage": f["soil_drainage"], "soil_awc_mm_per_m": f["soil_awc_mm_per_m"],
            "geological_structure_text": c["geological_structure_source_value"],
            "lineament_distance_m": f["lineament_dem_distance_m"], "lineament_density_1km": f["lineament_dem_density_1km_km_per_km2"],
            "gsi_lithology": f["gsi_lithology"], "gsi_group": f["gsi_group"], "gsi_formation": f["gsi_formation"], "gsi_age": f["gsi_age"],
            "gsi_distance_to_contact_m": f["gsi_distance_to_contact_m"], "gsi_fault_distance_m": f["distance_to_gsi_fault_m"],
            "gsi_nearest_fault_type": f["gsi_nearest_fault_type"], "gsi_nearest_fault_name": f["gsi_nearest_fault_name"],
            "gsi_thrust_distance_m": f["distance_to_gsi_thrust_m"], "gsi_fault_density_1km": f["gsi_fault_density_1km_km_per_km2"],
            "seasonality_text": c["seasonality_source_value"], "intervention_text": c["intervention_source_value"],
            "vulnerability_index_source": c["vulnerability_index_source_value"], "notes": f["processing_notes"],
        })
        row.update({c[0]: secondary.get(rid, {}).get(c[0], "") for c in SECONDARY_COLUMNS})
        if f["sg_clay_0_30cm_pct"]:
            row.update(soil_clay_pct=f["sg_clay_0_30cm_pct"], soil_silt_pct=f["sg_silt_0_30cm_pct"], soil_sand_pct=f["sg_sand_0_30cm_pct"],
                       soil_texture_from="SoilGrids 0-30 cm")
        elif f["soil_topsoil_clay_pct"]:
            row.update(soil_clay_pct=f["soil_topsoil_clay_pct"], soil_silt_pct=f["soil_topsoil_silt_pct"], soil_sand_pct=f["soil_topsoil_sand_pct"],
                       soil_texture_from="HWSD v2 0-20 cm")

        values, source = take(["lithology_source_text"])
        row.update(values, lithology_from=source)
        row["lithology_class"] = classify(values["lithology_source_text"], LITHOLOGY_CLASSES, "Mixed")
        if "ALLUVIUM" in f["gsi_group"].upper():
            row["gsi_lithology_class"] = "Alluvium / debris"
        else:
            row["gsi_lithology_class"] = classify(f["gsi_lithology"], LITHOLOGY_CLASSES, "Mixed")
        if row["lithology_class"]:
            row["rock_class"], row["rock_class_from"] = row["lithology_class"], "field record"
        elif row["gsi_lithology_class"]:
            row["rock_class"], row["rock_class_from"] = row["gsi_lithology_class"], "GSI map"
        values, source = take(["spring_type_source_text"])
        row.update(values, spring_type_from=source)
        row["spring_type_class"] = classify(values["spring_type_source_text"], SPRING_TYPE_CLASSES)
        values, source = take(["discharge_lpm_low", "discharge_lpm_high", "discharge_text"])
        row.update(values, discharge_from=source)
        if rid in field_flow:
            low, high, text = field_flow[rid]
            reference = f"; source report: {report_flow[rid]} L/min" if report_flow[rid] else ""
            row.update(discharge_lpm_low=low, discharge_lpm_high=high, discharge_text=text + reference,
                       discharge_from="DoCC field visit Oct 2026 (field form; the source report value is kept in discharge_text)")
        values, source = take(["ph", "tds_mg_l"])
        row.update(values, water_quality_from=source)
        values, source = take(["recharge_area_ha_low", "recharge_area_ha_high"])
        row.update(values, recharge_area_from=source)
        seasonal = watsan_seasonal.get((c["latitude_source_value"], c["longitude_source_value"])) if rid.startswith("WAT") else None
        if seasonal:
            row["discharge_dry_lpm"], row["discharge_monsoon_lpm"] = seasonal["Dry (LPM)"], seasonal["Monsoon (LPM)"]
        checks = [n for n in [row["notes"]] if n]
        if f["source_ph"] and not own[rid]["ph"]:
            checks.append(f"source pH {f['source_ph']} removed as invalid")
        if f["source_tds_mg_l"] and not own[rid]["tds_mg_l"]:
            checks.append(f"source TDS {f['source_tds_mg_l']} removed as invalid")
        if row["discharge_lpm_high"] and float(row["discharge_lpm_high"]) > 300:
            checks.append("discharge above 300 L/min: check units or record")
        if row["recharge_area_ha_high"] and float(row["recharge_area_ha_high"]) > 1000:
            checks.append("recharge area above 1000 ha: likely a watershed, not a springshed")
        if f["dem_catchment_window_edge"] == "Y":
            checks.append("DEM catchment runs beyond 20 km: point sits on a large river, drainage area is a minimum")
        if row["district"] and not f["district_clean"]:
            checks.append("district read from the WATSAN area name")
        row["notes"] = "; ".join(checks)

        key = {
            "elevation": row["elevation_m"], "slope": row["slope_deg"], "land cover": row["land_cover_point"],
            "rainfall": row["rainfall_mean_annual_mm"], "soil texture": row["soil_clay_pct"],
            "rock type": row["rock_class"], "fault distance": row["gsi_fault_distance_m"], "lineament distance": row["lineament_distance_m"],
            "spring type": row["spring_type_source_text"], "discharge": row["discharge_lpm_low"], "pH": row["ph"],
            "TDS": row["tds_mg_l"], "recharge area": row["recharge_area_ha_low"],
        }
        row["key_fields_filled"] = sum(1 for v in key.values() if v)
        missing = [k for k, v in key.items() if not v]
        if f["coordinate_status"] != "ok":
            mapped = {"elevation", "slope", "land cover", "rainfall", "soil texture", "lineament distance", "fault distance", "rock type"}
            reason = ["mapped values need a coordinate"] if mapped & set(missing) else []
            missing = [m for m in missing if m not in mapped] + reason
        elif not f["gsi_lithology"]:
            missing = [m + " (GSI layers cover Uttarakhand only)" if m in ("rock type", "fault distance") else m for m in missing]
        row["gaps_remaining"] = "; ".join(missing)
        rows.append(row)

    assign_split(rows, groups)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    # Clean files: one row per distinct spring (primary records only); all records stay in spring_master_v1.csv.
    for name, keep in (("clean", lambda r: True), ("input", lambda r: r["dataset_split"] == "input"),
                       ("test", lambda r: r["dataset_split"] == "test")):
        with (OUT_DIR / f"spring_master_v1_{name}.csv").open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=NAMES, lineterminator="\n")
            writer.writeheader()
            writer.writerows(r for r in rows if r["is_primary_record"] == "Y" and keep(r))
    with (OUT_DIR / "spring_master_v1.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=NAMES, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    with (OUT_DIR / "spring_master_v1_data_dictionary.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["column", "unit", "from", "meaning"])
        writer.writerows(COLUMNS)

    # Fill summary: per key column, filled counts by where the value came from.
    summary = [["column", "filled", "filled_pct", "from_this_record", "from_colocated_record", "blank"]]
    from_cols = {"lithology_source_text": "lithology_from", "spring_type_source_text": "spring_type_from",
                 "discharge_lpm_low": "discharge_from", "ph": "water_quality_from", "tds_mg_l": "water_quality_from",
                 "recharge_area_ha_low": "recharge_area_from"}
    for name in NAMES:
        filled = sum(1 for r in rows if r[name] not in ("", None))
        this = other = ""
        if name in from_cols:
            this = sum(1 for r in rows if r[name] and r[from_cols[name]] == "this record")
            other = sum(1 for r in rows if r[name] and r[from_cols[name]].startswith("colocated"))
        summary.append([name, filled, f"{100 * filled / len(rows):.0f}", this, other, len(rows) - filled])
    with (OUT_DIR / "spring_master_v1_fill_summary.csv").open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle, lineterminator="\n").writerows(summary)

    pair_rows = []
    by_id = {r["record_id"]: r for r in rows}
    for group, pairs in group_pairs.items():
        for a, b, decision, reason in pairs:
            ra, rb = by_id[a], by_id[b]
            pair_rows.append([group, decision, reason, a, ra["spring_name"], ra["village"], ra["elevation_source_m"],
                              b, rb["spring_name"], rb["village"], rb["elevation_source_m"], ra["elevation_m"]])
    pair_header = ["colocated_group_id", "decision", "reason", "record_a", "name_a", "village_a", "source_elevation_a_m",
                   "record_b", "name_b", "village_b", "source_elevation_b_m", "dem_elevation_m"]
    with (OUT_DIR / "spring_master_v1_duplicate_pairs.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(pair_header)
        writer.writerows(pair_rows)
    if not DECISIONS_CSV.exists():
        with DECISIONS_CSV.open("w", newline="", encoding="utf-8") as handle:
            csv.writer(handle, lineterminator="\n").writerow(DECISION_COLUMNS)

    write_workbook(rows, summary, pair_header, pair_rows)
    print(Counter(status_of[group[0]] for group in groups.values()), "groups;",
          Counter(d for pairs in group_pairs.values() for *_, d, _ in pairs), "pairs")
    primaries = sum(1 for r in rows if r["is_primary_record"] == "Y")
    print(f"{len(rows)} records, {primaries} primary (distinct springs)")
    print(Counter(r["key_fields_filled"] for r in rows if r["is_primary_record"] == "Y"))


GUIDE = [
    ("What this is", "One row per spring record from the common file (WATSAN, Spring Atlas annexure, Himmotthan, Atlas profiles, "
                     "Chamba report, Rathi 2021, Daniel 2021), with every gap filled that public data or a matching record can fill."),
    ("Input and test data", "The records are split in two. Input data is for building and tuning the tool. Test data (Tehri Garhwal, "
                            "the 10 Spring Atlas profile springs, and everything within 2 km of them) is kept aside: do not look at it or "
                            "tune rules on it, and score the finished tool on it once. Field-visit springs will be added to the test data. "
                            "The Input data and Test data sheets have one row per distinct spring (duplicates removed); the All records "
                            "sheet keeps every source record, including the duplicate copies, for tracing values back to their source. "
                            "Same in CSV: spring_master_v1_clean.csv (all 1,679 springs), _input.csv, _test.csv; spring_master_v1.csv has "
                            "all records."),
    ("Counting springs", "Some springs are listed by two sources. Filter is_primary_record = Y to keep one row per spring; "
                         "spring_id links the records of one spring."),
    ("Duplicates", "Records within 30 m of each other were compared by name, position word (Talla / Malla, I / II), village and "
                   "source elevation. duplicate_status is coloured: orange = same spring (merged), green = different springs that share "
                   "a coordinate (kept apart), yellow = needs review. The Duplicate review sheet lists only the pairs a person must "
                   "decide; enter same or different there and the decision goes into data/01_staged/common/"
                   "duplicate_review_decisions_v1.csv for the next build. Until then those records count as separate springs."),
    ("Mapped values", "Terrain, land cover, rainfall, soil, GSI geology, faults and lineaments come from the same layer for every spring, so springs can be "
                      "compared fairly. The source's own value is kept in a *_source column beside it."),
    ("Field values", "Discharge, pH, TDS, spring type, lithology and recharge area come from the record itself. Where it had none, a "
                     "record from another source within 30 m filled it; the *_from column names that record."),
    ("Not filled", "Nothing is estimated, averaged from neighbours or invented. Blank means no source has it; gaps_remaining lists what is missing."),
    ("Caveats", "Layer values describe the spring point or a circle around it, not the real recharge area. CHIRPS rainfall is ~5 km, "
                "HWSD soil ~1 km, SoilGrids 250 m, IMD ~25 km (2001-2025, 2012 missing). DEM lineaments are unvalidated candidates. "
                "Drainage area, wetness index and distance to stream come from a 30 m DEM: small hollows are invisible and a 30 m GPS "
                "error can move a spring into the next cell, so compare drainage_area_ha with drainage_area_max30m_ha. "
                "GSI geology is a state-scale map: a spring's polygon is a map overlay, not a field check, and it covers Uttarakhand only."),
    ("Rebuild", "python scripts/build_spring_features.py ... then python scripts/build_spring_master.py. "
                "See docs/data-gap-analysis.md."),
]


def write_workbook(rows, summary, pair_header, pair_rows):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Arial", size=10)
    numeric = re.compile(r"^-?\d+(\.\d+)?$")
    wb = Workbook()
    guide = wb.active
    guide.title = "Guide"
    guide.append(["Spring master v1", ""])
    guide["A1"].font = Font(name="Arial", bold=True, size=14)
    for item in GUIDE:
        guide.append(list(item))
    for row in guide.iter_rows(min_row=2):
        row[0].font = Font(name="Arial", bold=True, size=10)
        row[1].font = body_font
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 110

    def table(sheet, header, body, widths=None):
        sheet.append(header)
        for values in body:
            sheet.append([float(v) if isinstance(v, str) and numeric.match(v) else v for v in values])
        for cell in sheet[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        for i, name in enumerate(header, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = (widths or {}).get(name, 16)
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions

    wide = {n: 40 for n in ("lithology_source_text", "spring_type_source_text", "discharge_text", "gaps_remaining", "notes",
                            "geological_structure_text", "intervention_text", "land_use_source_text")}
    wide.update({n: 30 for n in ("lithology_from", "spring_type_from", "discharge_from", "water_quality_from", "recharge_area_from", "source_dataset")})
    wide.update({"duplicate_group": 60, "elevation_check": 45, "coordinate_correction": 40, "duplicate_status": 20})
    wide["split_reason"] = 35
    fills = {"same spring": "FCE4D6", "partly the same spring": "FCE4D6", "different springs": "E2EFDA", "needs review": "FFF2CC"}
    status_col, label_col = NAMES.index("duplicate_status") + 1, NAMES.index("duplicate_group") + 1
    sheets = (("Input data", lambda r: r["is_primary_record"] == "Y" and r["dataset_split"] == "input"),
              ("Test data", lambda r: r["is_primary_record"] == "Y" and r["dataset_split"] == "test"),
              ("All records", lambda r: True))
    for title, keep in sheets:
        part = [r for r in rows if keep(r)]
        sheet = wb.create_sheet(title)
        table(sheet, NAMES, [[r[n] for n in NAMES] for r in part], wide)
        for i, r in enumerate(part, start=2):
            if r["duplicate_status"]:
                fill = PatternFill("solid", fgColor=fills[r["duplicate_status"]])
                sheet.cell(row=i, column=status_col).fill = fill
                sheet.cell(row=i, column=label_col).fill = fill
    review = [p for p in pair_rows if p[1] == "review"]
    table(wb.create_sheet("Duplicate review"), pair_header + ["your_decision (same / different)", "note"],
          [p + ["", ""] for p in review],
          {"reason": 40, "name_a": 24, "name_b": 24, "your_decision (same / different)": 22, "note": 30})
    table(wb.create_sheet("Duplicate pairs"), pair_header, pair_rows, {"reason": 45, "name_a": 24, "name_b": 24})
    table(wb.create_sheet("Data dictionary"), ["column", "unit", "from", "meaning"], [list(c) for c in COLUMNS],
          {"column": 30, "from": 40, "meaning": 90})
    table(wb.create_sheet("Fill summary"), summary[0], summary[1:], {"column": 32})
    wb.save(OUT_DIR / "spring_master_v1.xlsx")


if __name__ == "__main__":
    main()
